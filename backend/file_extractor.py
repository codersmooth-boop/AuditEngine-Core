"""Extract text from PDF, XLSX, TXT files."""
import io
from pypdf import PdfReader
from openpyxl import load_workbook


def extract_text_from_file(filename: str, content: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            return _extract_pdf(content)
        if name.endswith(".xlsx") or name.endswith(".xls"):
            return _extract_xlsx(content)
        if name.endswith(".txt") or name.endswith(".csv") or name.endswith(".md"):
            return content.decode("utf-8", errors="ignore")
        # Fallback: try text decode
        return content.decode("utf-8", errors="ignore")[:50000]
    except Exception as e:
        return f"[extraction_error: {e}]"


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            parts.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)
